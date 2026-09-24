import pytest
from pathlib import Path
from docx import Document
from openpyxl import Workbook
from services.pdf_service import PDFService


def test_pdf_service_is_available():
    service = PDFService()
    assert service.is_available() is True


def test_pdf_service_convert_nonexistent_file(tmp_path):
    service = PDFService()
    nonexistent = tmp_path / "does_not_exist.docx"
    result = service.convert_to_pdf(nonexistent)
    assert result is None


def test_pdf_service_convert_docx_to_pdf(tmp_path):
    # Create test DOCX
    docx_path = tmp_path / "test_doc.docx"
    doc = Document()
    doc.add_heading("Test Heading", level=1)
    doc.add_paragraph("This is a test paragraph for PDF conversion.")
    doc.save(str(docx_path))

    service = PDFService()
    pdf_path = service.convert_to_pdf(docx_path, output_dir=tmp_path)

    assert pdf_path is not None
    assert pdf_path.exists()
    assert pdf_path.name == "test_doc.pdf"
    assert pdf_path.stat().st_size > 0


def test_pdf_service_convert_xlsx_to_pdf(tmp_path):
    # Create test XLSX
    xlsx_path = tmp_path / "test_sheet.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Test Sheet"
    ws.append(["No", "Uraian", "Jumlah"])
    ws.append([1, "Item A", 100000])
    wb.save(str(xlsx_path))

    service = PDFService()
    pdf_path = service.convert_to_pdf(xlsx_path, output_dir=tmp_path)

    assert pdf_path is not None
    assert pdf_path.exists()
    assert pdf_path.name == "test_sheet.pdf"
    assert pdf_path.stat().st_size > 0
